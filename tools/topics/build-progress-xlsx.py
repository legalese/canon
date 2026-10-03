#!/usr/bin/env python3
"""Build encoding-progress.xlsx: every topic Act, grouped by jurisdiction, with a status column
for each stage of the encoding pipeline."""
import glob, importlib.util, json, os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
_s3 = importlib.util.spec_from_file_location("scaffold3", os.path.join(HERE, "scaffold-round3.py"))
scaffold3 = importlib.util.module_from_spec(_s3); _s3.loader.exec_module(scaffold3)
spec = json.load(open(os.path.join(HERE, "topic-acts.json"), encoding="utf-8"))
ver = json.load(open(os.path.join(HERE, "verification.json"), encoding="utf-8"))
for _extra in (os.path.join(HERE, "round2", "verification.json"),
               os.path.join(HERE, "round3", "verification.json"),
               os.path.join(HERE, "round4", "verification.json")):
    if os.path.exists(_extra):
        ver += json.load(open(_extra, encoding="utf-8"))

ORDER = [("sg", "Singapore"), ("au-cth", "AU-Commonwealth"), ("nz", "New Zealand"),
         ("au-wa", "AU-WA"), ("au-sa", "AU-SA"), ("au-nsw", "AU-NSW"), ("au-vic", "AU-Vic"),
         ("au-qld", "AU-Qld"), ("au-tas", "AU-Tas"), ("au-act", "AU-ACT"), ("au-nt", "AU-NT"),
         ("uk", "UK"), ("us-federal", "US-Federal"), ("us-ca", "US-California"),
         ("us-ny", "US-New York"), ("us-tx", "US-Texas"), ("ca-federal", "CA-Federal"),
         ("ca-on", "CA-Ontario"), ("ca-bc", "CA-British Columbia"), ("in-central", "India"), ("ie", "Ireland"),
         ("za", "South Africa"), ("hk", "Hong Kong"), ("il", "Israel")]
# round 3: the remaining US states and DC, then the remaining Canadian provinces and territories
ORDER += [("us-" + c.lower(), "US-" + n) for n, c in sorted(scaffold3.US.items())
          if "us-" + c.lower() not in dict(ORDER)]
ORDER += [("ca-" + c.lower(), "CA-" + n) for n, c in sorted(scaffold3.CA.items())
          if "ca-" + c.lower() not in dict(ORDER)]
STAGES = ["Indexed", "Source text", "AI pass #1", "Human pass #1", "AI pass #2", "Human pass #2"]
ROOT = os.path.dirname(os.path.dirname(HERE))
WORKER = json.load(open(os.path.join(ROOT, "tools", "worker", "config.json"), encoding="utf-8"))
DIRS = {"sg": "singapore", "au-cth": "commonwealth-of-australia", "nz": "new-zealand",
        "au-wa": "western-australia", "au-sa": "south-australia", "au-nsw": "new-south-wales",
        "au-vic": "victoria", "au-qld": "queensland", "au-tas": "tasmania", "au-act": "australian-capital-territory",
        "au-nt": "northern-territory", "uk": "united-kingdom", "us-federal": "united-states",
        "us-ca": "california", "us-ny": "new-york", "us-tx": "texas", "ca-federal": "canada",
        "ca-on": "ontario", "ca-bc": "british-columbia", "in-central": "india", "ie": "ireland",
        "za": "south-africa", "hk": "hong-kong", "il": "israel"}
DIRS.update({"us-" + c.lower(): scaffold3.slug(n) for n, c in scaffold3.US.items()})
DIRS.update({"ca-" + c.lower(): scaffold3.slug(n) for n, c in scaffold3.CA.items()})
SUBJECT = {}
for _sj in glob.glob(os.path.join(ROOT, "subjects", "*", "*", "subject.json")):
    try:
        _name = json.load(open(_sj, encoding="utf-8")).get("display_name")
    except Exception:
        continue
    SUBJECT[(os.path.basename(os.path.dirname(os.path.dirname(_sj))), _name)] = os.path.dirname(_sj)
FETCH_WHY = {"not-built": None, "ready": None,
             "needs-api-key": "Register sits behind a bot check; bulk access needs the PCO API key.",
             "script-blocked": "Register refuses scripted access; not circumvented.",
             "truncated-response": "Register cuts every download off at ~32 KB.",
             "have-text": None}


def source_text(jur, act):
    """(status, comment) for the Source text column, read from the subject directory."""
    d = SUBJECT.get((DIRS[jur], act))
    if d is None:
        return "", None
    bundle = os.path.join(d, "registers", "source-bundle")
    texts = [p for p in glob.glob(os.path.join(bundle, "*")) if p.endswith((".txt", ".xml", ".html"))]
    texts += glob.glob(os.path.join(d, "source*.txt"))      # e.g. israel/...: source-law-he.txt
    if texts:
        note = []
        for mp in glob.glob(os.path.join(bundle, "*.meta.json")):
            m = json.load(open(mp, encoding="utf-8"))
            if m.get("version"):
                note.append("%s (%s), fetched %s." % (m["version"], m.get("version_kind", "version"),
                                                     m.get("retrieved_date", "?")))
            if m.get("text_caveat"):
                note.append("CAVEAT: " + m["text_caveat"])
        return "Done", " ".join(note) or None
    if glob.glob(os.path.join(d, "*.l4")):
        return "", None
    pol = WORKER["jurisdictions"].get(DIRS[jur], {})
    lic = pol.get("licence")
    if lic == "unclear":
        return "", ("Licence decision pending: see tools/topics/round3/licence-review.md. "
                    "Scaffolded metadata-only, so no statute text may be deposited yet.")
    if lic not in ("open", "permission"):
        return "Blocked", "Licence: %s. No statute text until a human clears it." % lic
    why = FETCH_WHY.get(pol.get("fetch"))
    if pol.get("fetch", "").startswith("window:"):
        why = "SSO allows scripted fetching only 03:00-07:00 SGT; the nightly harvest fetches it."
        return "", why
    return ("Blocked", why) if why else ("", None)


TOPIC_RANK = {t: i for i, t in enumerate(spec["topics"])}
TOPIC_RANK["Other"] = len(TOPIC_RANK)

# Acts already encoded that fall outside the 12 topics. Listed so the tracker shows all
# encoding work, not only work on the topic list.
EXTRA = [
    {"jurisdiction": "sg", "act": "Penal Code 1871", "topics": ["Other"], "status": "VERIFIED"},
    {"jurisdiction": "au-wa", "act": "Dog Act 1976", "topics": ["Other"], "status": "VERIFIED"},
    {"jurisdiction": "il", "act": "Regulation of Work on Refrigeration or Air-Conditioning Systems Law, 5785-2025",
     "topics": ["Other"], "status": "VERIFIED"},
]

# Encoding work that already exists for Acts on this list.
EXISTING = {
    ("il", "Regulation of Work on Refrigeration or Air-Conditioning Systems Law, 5785-2025"): (
        {"AI pass #1": "Done"},
        "canon main: subjects/israel/refrigeration-or-air-conditioning-work-law-2025, v0.1.0, status "
        "draft. All 63 sections and 4 schedules mapped; checked on l4 unstable-20260907. No HG1."),
    ("sg", "Penal Code 1871"): (
        {"AI pass #1": "Done"},
        "canon main: subjects/singapore/penal-code-1871, v0.9.0, status draft. All 525 live "
        "sections, ss 1 to 512 (Aswathy Satheesan, completed 2026-09-16). Two fidelity "
        "verification passes recorded in registers/. No HG1."),
    ("au-wa", "Dog Act 1976"): (
        {"AI pass #1": "Done"},
        "Whole Act: 18 modules, 654 assertions, v0.2.0, status draft. Authoritative copy in "
        "legalese/l4-ide (etc/go/subjects/dog-act-1976); the copy in the local canon clone is "
        "uncommitted. No HG1 or HG2."),
    ("au-wa", "Residential Tenancies Act 1987"): (
        {"AI pass #1": "Done"},
        "canon: subjects/western-australia/residential-tenancies-act, v0.6.0, status draft. "
        "16 modules, operational core encoded; see its NOTES.md for what is left out."),
    ("uk", "Housing Act 1988"): (
        {"AI pass #1": "In progress"},
        "Schedule 2 (grounds for possession) only: OPC UK/HA1988_Sch2_*.l4, and the "
        "housing-act-wizard deployment covers Grounds 8, 10 and 11."),
}

A = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=A, size=10, bold=True, color="FFFFFF")
BODY = Font(name=A, size=10)
BOLD = Font(name=A, size=10, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BAND = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3FA")]
CENTER = Alignment(horizontal="center", vertical="top")
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
ws = wb.active
ws.title = "Progress"
hdr = ["Jurisdiction", "Act", "Category"] + STAGES
ws.append(hdr)
for c in range(1, len(hdr) + 1):
    cell = ws.cell(1, c)
    cell.font = H_FONT; cell.fill = H_FILL; cell.border = BOX
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if c > 3 else "left")
ws.row_dimensions[1].height = 30

rank = {j: k for k, (j, _) in enumerate(ORDER)}
label = dict(ORDER)
rows = sorted(ver + EXTRA, key=lambda r: (rank[r["jurisdiction"]],
                                  min(TOPIC_RANK[t] for t in r["topics"]), r["act"]))
band = -1
prev = None
for r in rows:
    jur, act = r["jurisdiction"], r["act"]
    if jur != prev:
        band += 1; prev = jur
    topics = sorted(r["topics"], key=lambda t: TOPIC_RANK[t])
    stages = {s: "" for s in STAGES}
    stages["Indexed"] = "Done"
    stages["Source text"], src_note = source_text(jur, act)
    comment = None
    if (jur, act) in EXISTING:
        upd, comment = EXISTING[(jur, act)]
        stages.update(upd)
    ws.append([label[jur], act, topics[0]] + [stages[s] for s in STAGES])
    n = ws.max_row
    fill = BAND[band % 2]
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(n, c)
        cell.font = BOLD if c == 1 else BODY
        cell.border = BOX; cell.fill = fill
        cell.alignment = CENTER if c > 3 else WRAP
    if len(topics) > 1:
        ws.cell(n, 3).comment = Comment("Also covers: " + "; ".join(topics[1:]), "canon")
    if topics[0] == "Other":
        ws.cell(n, 3).comment = Comment("Outside the 12 topics. Listed because it is already "
                                        "encoded.", "canon")
    if r["status"] == "UNVERIFIED":
        ws.cell(n, 2).comment = Comment("In-force status UNVERIFIED (2026-09-22): probably not yet "
                                        "in force; commencement not found.", "canon")
    if r["status"] == "NOT YET IN FORCE":
        ws.cell(n, 2).comment = Comment("Enacted but NOT YET IN FORCE (checked 2026-09-22). "
                                        "Kept deliberately: encoding ahead of commencement.", "canon")
    if comment:
        ws.cell(n, hdr.index("AI pass #1") + 1).comment = Comment(comment, "canon")
    if src_note:
        ws.cell(n, hdr.index("Source text") + 1).comment = Comment(src_note, "canon")
last = ws.max_row

# Stage columns take a fixed vocabulary, colour-coded.
dv = DataValidation(type="list", formula1='"Not started,In progress,Done,Blocked"', allow_blank=True)
ws.add_data_validation(dv)
rng = "D2:%s%d" % (get_column_letter(len(hdr)), last)
dv.add(rng)
for val, colour in (("Done", "C6EFCE"), ("In progress", "FFEB9C"), ("Blocked", "FFC7CE")):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"%s"' % val],
                                                  fill=PatternFill("solid", fgColor=colour)))
for c, w in zip("ABCDEFGHI", (17, 62, 24, 11, 11, 11, 13, 11, 13)):
    ws.column_dimensions[c].width = w
ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(hdr)), last)

# ------------------------------------------------------------------ Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "Stages done, by jurisdiction"; sm["A1"].font = Font(name=A, size=14, bold=True)
sm["A2"] = ("Counts of 'Done' in each stage column of the Progress sheet. Stage cells accept "
            "Not started, In progress, Done or Blocked.")
sm["A2"].font = Font(name=A, size=9, italic=True, color="808080")
head = ["Jurisdiction", "Acts"] + STAGES
for c, h in enumerate(head, 1):
    cell = sm.cell(4, c, h); cell.font = H_FONT; cell.fill = H_FILL; cell.border = BOX
    cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
for i, (_, lab) in enumerate(ORDER, 5):
    sm.cell(i, 1, lab).font = BOLD
    sm.cell(i, 2, "=COUNTIF(Progress!$A$2:$A$%d,A%d)" % (last, i))
    for k in range(len(STAGES)):
        col = get_column_letter(4 + k)
        sm.cell(i, 3 + k, '=COUNTIFS(Progress!$A$2:$A$%d,$A%d,Progress!$%s$2:$%s$%d,"Done")'
                % (last, i, col, col, last))
    for c in range(1, len(head) + 1):
        sm.cell(i, c).border = BOX
        if c > 1:
            sm.cell(i, c).font = BODY; sm.cell(i, c).alignment = Alignment(horizontal="center")
t = 5 + len(ORDER)
sm.cell(t, 1, "Total").font = BOLD
for c in range(2, len(head) + 1):
    col = get_column_letter(c)
    cell = sm.cell(t, c, "=SUM(%s5:%s%d)" % (col, col, t - 1))
    cell.font = BOLD; cell.border = BOX; cell.alignment = Alignment(horizontal="center")
sm.cell(t, 1).border = BOX
for c, w in zip("ABCDEFGH", (18, 8, 10, 11, 11, 13, 11, 13)):
    sm.column_dimensions[c].width = w

out = os.path.join(HERE, "encoding-progress.xlsx")
wb.save(out)
print("saved %s, %d Acts" % (out, last - 1))
