#!/usr/bin/env python3
"""Build topic-acts.xlsx from topic-acts.json and verification.json."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
spec = json.load(open(os.path.join(HERE, "topic-acts.json"), encoding="utf-8"))
ver = json.load(open(os.path.join(HERE, "verification.json"), encoding="utf-8"))
browser = json.load(open(os.path.join(HERE, "browser-verified.json"), encoding="utf-8"))

ORDER = [("sg", "Singapore"), ("au-cth", "AU-Commonwealth"), ("nz", "New Zealand"),
         ("au-wa", "AU-WA"), ("au-sa", "AU-SA"), ("au-nsw", "AU-NSW"), ("au-vic", "AU-Vic"),
         ("au-qld", "AU-Qld"), ("au-tas", "AU-Tas"), ("uk", "UK")]
LABEL = dict(ORDER)
SOURCE = {
    "sg": "SSO Act manifest (tools/sso-acts-manifest.json)",
    "au-cth": "Federal Register of Legislation API, principal Acts in force",
    "nz": "legislation.govt.nz title search, status badge (browser)",
    "au-wa": "WA subjects in canon (legislation.wa.gov.au)",
    "au-sa": "SA Current Titles list, 11 Sep 2026",
    "au-nsw": "NSW Public Acts in force table (browser)",
    "au-vic": "Victorian Legislation in-force search",
    "au-qld": "OQPC browse data source, in force",
    "au-tas": "EnAct browse data source, in force",
    "uk": "legislation.gov.uk data feed; exists and not marked repealed (weaker check)",
}
topics = spec["topics"]
cth_only = set(spec["commonwealth_matters_for_states"])
notes = spec["notes"]

A = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=A, size=10, bold=True, color="FFFFFF")
BODY = Font(name=A, size=10)
BOLD = Font(name=A, size=10, bold=True)
GREY = Font(name=A, size=9, italic=True, color="808080")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AMBER = PatternFill("solid", fgColor="FFF2CC")
PALE = PatternFill("solid", fgColor="F2F2F2")

status = {(r["jurisdiction"], r["act"]): r["status"] for r in ver}
wb = Workbook()

# ------------------------------------------------------------------ Read me
ws = wb.active
ws.title = "Read me"
lines = [
    ("Acts by subject matter", None, Font(name=A, size=14, bold=True)),
    ("Twelve topics x ten jurisdictions, for choosing which Acts to encode", None, Font(name=A, size=11, italic=True)),
    (None, None, None),
    ("Built", "2026-09-22", BOLD),
    ("Method", "Each topic was mapped to the Acts that govern it in each jurisdiction, then every Act was "
               "checked against that jurisdiction's official register for its exact current title and that "
               "it is in force. Titles that failed the check were corrected or dropped; the corrections are "
               "recorded in the Notes column of the Acts sheet.", BOLD),
    ("Scope", "Acts only. Where a topic is governed by subordinate legislation or retained EU law rather "
              "than an Act (UK holiday entitlement; UK flight-delay compensation), that is noted rather than "
              "filled.", BOLD),
    ("Commonwealth matters", "In Australia, pensions and social security, government benefits, personal "
              "income tax, passenger compensation and immigration are federal. State cells for those topics "
              "read 'Commonwealth matter' and are genuinely empty, not missing.", BOLD),
    ("Status", "VERIFIED: title matches the official register and the Act is in force. NOT YET IN FORCE: "
               "enacted, awaiting commencement, kept deliberately. UK VERIFIED is a weaker check than the "
               "others: the Act exists on legislation.gov.uk and is not marked repealed.", BOLD),
    ("Not a ranking", "The mapping is a judgement about which Acts govern each topic, checked for accuracy. "
                      "It is not a measure of how often each Act is used.", BOLD),
]
for i, (a, b, f) in enumerate(lines, 1):
    ws.cell(i, 1, a).font = f or BODY
    ws.cell(i, 2, b).font = BODY
    ws.cell(i, 1).alignment = Alignment(vertical="top")
    ws.cell(i, 2).alignment = WRAP
    if b and len(b) > 90:
        ws.row_dimensions[i].height = 15 * (len(b) // 95 + 1)
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 110

# ------------------------------------------------------------------ Matrix
ws = wb.create_sheet("Matrix")
ws.cell(1, 1, "Jurisdiction")
for j, t in enumerate(topics, 2):
    ws.cell(1, j, t)
for c in range(1, len(topics) + 2):
    ws.cell(1, c).font = H_FONT; ws.cell(1, c).fill = H_FILL
    ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.cell(1, c).border = BOX
for i, (jur, label) in enumerate(ORDER, 2):
    ws.cell(i, 1, label).font = BOLD
    ws.cell(i, 1).alignment = Alignment(vertical="top")
    ws.cell(i, 1).border = BOX
    for j, t in enumerate(topics, 2):
        acts = spec["map"].get(jur, {}).get(t, [])
        cell = ws.cell(i, j)
        if acts:
            cell.value = "\n".join(a + (" (not yet in force)" if status.get((jur, a)) == "NOT YET IN FORCE" else "")
                                   for a in acts)
            cell.font = BODY
            if any(status.get((jur, a)) == "NOT YET IN FORCE" for a in acts):
                cell.fill = AMBER
        elif jur.startswith("au-") and jur != "au-cth" and t in cth_only:
            cell.value = "Commonwealth matter"; cell.font = GREY; cell.fill = PALE
        else:
            cell.value = notes.get("%s|%s" % (jur, t), "no Act identified")
            cell.font = GREY; cell.fill = PALE
        cell.alignment = WRAP; cell.border = BOX
    ws.row_dimensions[i].height = 110
ws.column_dimensions["A"].width = 17
for j in range(2, len(topics) + 2):
    ws.column_dimensions[get_column_letter(j)].width = 24
ws.freeze_panes = "B2"

# ------------------------------------------------------------------ Acts
ws = wb.create_sheet("Acts")
hdr = ["Jurisdiction", "Act", "Topics covered", "Topic count", "Status", "Checked against", "Notes"]
ws.append(hdr)
for c in range(1, len(hdr) + 1):
    ws.cell(1, c).font = H_FONT; ws.cell(1, c).fill = H_FILL; ws.cell(1, c).border = BOX
    ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center")
rank = {j: k for k, (j, _) in enumerate(ORDER)}
rows = sorted(ver, key=lambda r: (rank[r["jurisdiction"]], -len(r["topics"]), r["act"]))
for r in rows:
    jur = r["jurisdiction"]
    note = "; ".join(v for k, v in notes.items()
                     if k.split("|")[0] == jur and k.split("|")[1] in r["topics"] and r["act"] in v)
    bnote = browser.get(jur, {}).get(r["act"], "")
    if " -- " in bnote:
        note = (note + "; " if note else "") + bnote.split(" -- ", 1)[1]
    ws.append([LABEL[jur], r["act"], "; ".join(r["topics"]), None, r["status"], SOURCE[jur], note])
    n = ws.max_row
    ws.cell(n, 4, '=IF(C%d="",0,LEN(C%d)-LEN(SUBSTITUTE(C%d,";",""))+1)' % (n, n, n))
    for c in range(1, len(hdr) + 1):
        ws.cell(n, c).font = BODY; ws.cell(n, c).alignment = WRAP; ws.cell(n, c).border = BOX
    if r["status"] == "NOT YET IN FORCE":
        for c in range(1, len(hdr) + 1):
            ws.cell(n, c).fill = AMBER
last = ws.max_row
for c, w in zip("ABCDEFG", (16, 58, 46, 11, 18, 44, 60)):
    ws.column_dimensions[c].width = w
ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:G%d" % last

# ------------------------------------------------------------------ Summary
ws = wb.create_sheet("Summary")
ws["A1"] = "Acts per jurisdiction"; ws["A1"].font = Font(name=A, size=14, bold=True)
for c, h in enumerate(["Jurisdiction", "Distinct Acts", "Verified", "Not yet in force", "Topics with an Act"], 1):
    cell = ws.cell(3, c, h); cell.font = H_FONT; cell.fill = H_FILL; cell.border = BOX
    cell.alignment = Alignment(wrap_text=True, vertical="center")
for i, (jur, label) in enumerate(ORDER, 4):
    ws.cell(i, 1, label).font = BOLD
    ws.cell(i, 2, '=COUNTIF(Acts!$A$2:$A$%d,A%d)' % (last, i))
    ws.cell(i, 3, '=COUNTIFS(Acts!$A$2:$A$%d,A%d,Acts!$E$2:$E$%d,"VERIFIED")' % (last, i, last))
    ws.cell(i, 4, '=COUNTIFS(Acts!$A$2:$A$%d,A%d,Acts!$E$2:$E$%d,"NOT YET IN FORCE")' % (last, i, last))
    ws.cell(i, 5, sum(1 for t in topics if spec["map"].get(jur, {}).get(t)))
    for c in range(1, 6):
        ws.cell(i, c).border = BOX
        if c > 1: ws.cell(i, c).font = BODY
t = 4 + len(ORDER)
ws.cell(t, 1, "Total").font = BOLD
for c in (2, 3, 4):
    col = get_column_letter(c)
    ws.cell(t, c, "=SUM(%s4:%s%d)" % (col, col, t - 1)).font = BOLD
for c, w in zip("ABCDE", (20, 14, 12, 16, 18)):
    ws.column_dimensions[c].width = w
ws.cell(t + 2, 1, "Topics with an Act counts Commonwealth-matter cells as empty for the states, "
                  "which is correct: those topics have no state Act.").font = GREY

out = os.path.join(HERE, "topic-acts.xlsx")
wb.save(out)
print("saved", out, "rows", last - 1)
